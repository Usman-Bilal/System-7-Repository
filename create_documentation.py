"""
Create comprehensive Word documentation for System 7 Repository
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def add_heading_custom(doc, text, level=1):
    """Add a custom styled heading"""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return heading

def add_code_block(doc, code, language="python"):
    """Add a code block with formatting"""
    para = doc.add_paragraph()
    para.style = 'No Spacing'
    run = para.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return para

def create_documentation():
    """Create the main documentation"""
    doc = Document()
    
    # Title page
    title = doc.add_heading('System 7 Repository', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Week 1 - Python Programming Tasks')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(0, 102, 204)
    
    author = doc.add_paragraph('Usman Bilal')
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    date = doc.add_paragraph('October 2026')
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # Table of Contents
    add_heading_custom(doc, 'Table of Contents', 1)
    toc_items = [
        ('1. Introduction', '3'),
        ('2. Day 1: CSV Data Analysis', '4'),
        ('3. Day 2: Calculator & Guessing Game', '6'),
        ('4. Day 3: Student Record Manager', '8'),
        ('5. Day 4: OOP - Bank & Library Systems', '10'),
        ('6. Repository Structure', '12'),
        ('7. How to Run', '13'),
    ]
    
    for item, page in toc_items:
        para = doc.add_paragraph(f'{item}', style='List Number')
    
    doc.add_page_break()
    
    # Introduction
    add_heading_custom(doc, '1. Introduction', 1)
    doc.add_paragraph(
        'This document contains comprehensive documentation for the System 7 Repository, '
        'which includes Python programming tasks completed during Week 1 of the internship program. '
        'Each day focuses on different programming concepts and skills, building from basic data analysis '
        'to advanced object-oriented programming.'
    )
    
    doc.add_paragraph('GitHub Repository:', style='List Bullet')
    doc.add_paragraph('https://github.com/Usman-Bilal/System-7-Repository')
    
    add_heading_custom(doc, 'Skills Covered', 2)
    skills = [
        'Data Analysis with Pandas, Matplotlib, Seaborn',
        'Input/Output and Control Flow',
        'File Handling and JSON',
        'Object-Oriented Programming (OOP)',
        'Error Handling and Validation',
        'Git and GitHub Workflow'
    ]
    for skill in skills:
        doc.add_paragraph(skill, style='List Bullet')
    
    doc.add_page_break()
    
    # DAY 1
    add_heading_custom(doc, '2. Day 1: CSV Data Analysis', 1)
    
    add_heading_custom(doc, 'Task Overview', 2)
    doc.add_paragraph(
        'Analyze a CSV dataset and produce charts and summary statistics using Pandas, Matplotlib, and Seaborn.'
    )
    
    add_heading_custom(doc, 'Skills Used', 2)
    day1_skills = ['Pandas: Data manipulation and analysis', 'Matplotlib: Data visualization', 'Seaborn: Statistical plots']
    for skill in day1_skills:
        doc.add_paragraph(skill, style='List Bullet')
    
    add_heading_custom(doc, 'Files Created', 2)
    day1_files = [
        'sample_data.csv - Employee dataset (20 records)',
        'data_analysis.py - Complete Python analysis script',
        'analysis_notebook.ipynb - Interactive Jupyter notebook',
        'analysis_charts.png - 9-panel visualization grid',
        'detailed_analysis.png - 4 detailed statistical plots'
    ]
    for file in day1_files:
        doc.add_paragraph(file, style='List Bullet')
    
    add_heading_custom(doc, 'Key Findings', 2)
    findings = [
        'Total employees analyzed: 20',
        'Average salary: $75,550',
        'Salary range: $58,000 - $102,000',
        'Strong correlation between experience and salary (0.888)',
        'Engineering department has highest average salary',
        'Average performance score: 8.5/10'
    ]
    for finding in findings:
        doc.add_paragraph(finding, style='List Bullet')
    
    add_heading_custom(doc, 'Sample Code', 2)
    code_sample = '''import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

# Load dataset
df = pd.read_csv('sample_data.csv')

# Display summary statistics
print(df.describe())

# Create visualization
plt.figure(figsize=(10, 6))
plt.hist(df['Salary'], bins=15, color='skyblue')
plt.title('Salary Distribution')
plt.xlabel('Salary ($)')
plt.ylabel('Frequency')
plt.show()'''
    
    add_code_block(doc, code_sample)
    
    doc.add_page_break()
    
    # DAY 2
    add_heading_custom(doc, '3. Day 2: Calculator & Number Guessing Game', 1)
    
    add_heading_custom(doc, 'Task Overview', 2)
    doc.add_paragraph(
        'Build a command-line calculator and number-guessing game using input/output, loops, and functions.'
    )
    
    add_heading_custom(doc, 'Calculator Features', 2)
    calc_features = [
        'Basic operations: Addition, Subtraction, Multiplication, Division',
        'Advanced operations: Power, Modulus',
        'Input validation and error handling',
        'User-friendly menu interface',
        'Continuous operation until user exits'
    ]
    for feature in calc_features:
        doc.add_paragraph(feature, style='List Bullet')
    
    add_heading_custom(doc, 'Guessing Game Features', 2)
    game_features = [
        'Classic mode: Guess number between 1-100 in 10 attempts',
        'Three difficulty levels (Easy, Medium, Hard)',
        'Hint system (too high/too low)',
        'Attempt tracking',
        'Play again functionality'
    ]
    for feature in game_features:
        doc.add_paragraph(feature, style='List Bullet')
    
    add_heading_custom(doc, 'Sample Calculator Code', 2)
    calc_code = '''def add(a, b):
    return a + b

def divide(a, b):
    if b == 0:
        return "Error: Cannot divide by zero"
    return a / b

# Main calculator loop
while True:
    display_menu()
    choice = input("Select operation (1-7): ")
    
    if choice == '7':
        print("Thank you for using the calculator!")
        break
    
    num1 = float(input("Enter first number: "))
    num2 = float(input("Enter second number: "))
    
    if choice == '1':
        result = add(num1, num2)
        print(f"Result: {num1} + {num2} = {result}")'''
    
    add_code_block(doc, calc_code)
    
    add_heading_custom(doc, 'Sample Guessing Game Code', 2)
    game_code = '''import random

secret_number = random.randint(1, 100)
attempts = 0
max_attempts = 10

while attempts < max_attempts:
    attempts += 1
    guess = int(input(f"Attempt {attempts}/{max_attempts} - Enter your guess: "))
    
    if guess == secret_number:
        print(f"CONGRATULATIONS! You guessed it in {attempts} attempts!")
        break
    elif guess < secret_number:
        print("Too low! Try higher.")
    else:
        print("Too high! Try lower.")'''
    
    add_code_block(doc, game_code)
    
    doc.add_page_break()
    
    # DAY 3
    add_heading_custom(doc, '4. Day 3: Student Record Manager', 1)
    
    add_heading_custom(doc, 'Task Overview', 2)
    doc.add_paragraph(
        'Create a student record manager that saves to a file using dictionaries, file handling, and exceptions.'
    )
    
    add_heading_custom(doc, 'Features', 2)
    day3_features = [
        'Add new student records (ID, Name, Age, Grade, Email)',
        'View all student records in formatted table',
        'Search student by ID',
        'Update existing student information',
        'Delete student records',
        'Persistent storage using JSON file',
        'Data validation (email format, grade validation)',
        'Duplicate ID prevention'
    ]
    for feature in day3_features:
        doc.add_paragraph(feature, style='List Bullet')
    
    add_heading_custom(doc, 'Sample Code', 2)
    student_code = '''import json

def save_students(students):
    """Save student records to JSON file"""
    with open('students.json', 'w') as file:
        json.dump(students, file, indent=4)

def add_student(students):
    """Add a new student"""
    student_id = input("Enter Student ID: ")
    
    if student_id in students:
        print("Student ID already exists!")
        return
    
    name = input("Enter Student Name: ")
    age = int(input("Enter Age: "))
    grade = input("Enter Grade (A/B/C/D/F): ").upper()
    email = input("Enter Email: ")
    
    students[student_id] = {
        'name': name,
        'age': age,
        'grade': grade,
        'email': email
    }
    
    save_students(students)
    print(f"Student '{name}' added successfully!")'''
    
    add_code_block(doc, student_code)
    
    doc.add_page_break()
    
    # DAY 4
    add_heading_custom(doc, '5. Day 4: OOP - Bank Account & Library Systems', 1)
    
    add_heading_custom(doc, 'Task Overview', 2)
    doc.add_paragraph(
        'Design Bank Account and Library management systems using Object-Oriented Programming principles.'
    )
    
    add_heading_custom(doc, 'Bank Account System', 2)
    bank_features = [
        'Account creation with unique account numbers',
        'Deposit and withdrawal operations',
        'Transaction history tracking',
        'Interest calculation (Savings accounts)',
        'Overdraft protection (Checking accounts)',
        'Inheritance: SavingsAccount and CheckingAccount classes'
    ]
    for feature in bank_features:
        doc.add_paragraph(feature, style='List Bullet')
    
    add_heading_custom(doc, 'Library Management System', 2)
    library_features = [
        'Book catalog management (add, remove, search)',
        'Member registration system',
        'Book borrowing and returning',
        'Due date tracking (14 days)',
        'Fine calculation for overdue books ($0.50/day)',
        'Book availability checking',
        'Search by title, author, or ISBN'
    ]
    for feature in library_features:
        doc.add_paragraph(feature, style='List Bullet')
    
    add_heading_custom(doc, 'OOP Concepts Implemented', 2)
    oop_concepts = [
        'Classes and Objects',
        'Encapsulation (private attributes)',
        'Inheritance (parent/child classes)',
        'Methods (instance methods)',
        'Constructor (__init__)',
        'String representation (__str__)'
    ]
    for concept in oop_concepts:
        doc.add_paragraph(concept, style='List Bullet')
    
    add_heading_custom(doc, 'Sample Bank Account Code', 2)
    bank_code = '''class BankAccount:
    """Base Bank Account class"""
    
    account_counter = 1000
    
    def __init__(self, account_holder, initial_balance=0):
        self.account_number = BankAccount.account_counter
        BankAccount.account_counter += 1
        self.account_holder = account_holder
        self._balance = initial_balance
        self.transactions = []
    
    def deposit(self, amount):
        """Deposit money into account"""
        if amount <= 0:
            print("Deposit amount must be positive!")
            return False
        
        self._balance += amount
        print(f"${amount:.2f} deposited successfully!")
        return True
    
    def get_balance(self):
        """Get current balance"""
        return self._balance

class SavingsAccount(BankAccount):
    """Savings account with interest"""
    
    def __init__(self, holder, balance=0, interest_rate=0.03):
        super().__init__(holder, balance)
        self.interest_rate = interest_rate
    
    def add_interest(self):
        """Add interest to account"""
        interest = self._balance * self.interest_rate
        self._balance += interest
        print(f"Interest of ${interest:.2f} added!")'''
    
    add_code_block(doc, bank_code)
    
    doc.add_page_break()
    
    # Repository Structure
    add_heading_custom(doc, '6. Repository Structure', 1)
    
    structure = '''System-7-Repository/
├── README.md
├── requirements.txt
├── .gitignore
├── Day1/
│   ├── README.md
│   ├── sample_data.csv
│   ├── data_analysis.py
│   ├── analysis_notebook.ipynb
│   ├── analysis_charts.png
│   └── detailed_analysis.png
├── Day2/
│   ├── README.md
│   ├── calculator.py
│   ├── number_guessing_game.py
│   └── PLAY_INSTRUCTIONS.txt
├── Day3/
│   ├── README.md
│   ├── student_manager.py
│   └── students.json (auto-generated)
└── Day4/
    ├── README.md
    ├── bank_account.py
    ├── bank_demo.py
    ├── library_system.py
    └── library_demo.py'''
    
    add_code_block(doc, structure, "text")
    
    doc.add_page_break()
    
    # How to Run
    add_heading_custom(doc, '7. How to Run', 1)
    
    add_heading_custom(doc, 'Prerequisites', 2)
    doc.add_paragraph('Python 3.7 or higher installed')
    doc.add_paragraph('Git installed')
    
    add_heading_custom(doc, 'Installation Steps', 2)
    install_steps = [
        'Clone the repository:',
        '   git clone https://github.com/Usman-Bilal/System-7-Repository.git',
        '',
        'Navigate to the repository:',
        '   cd System-7-Repository',
        '',
        'Install required packages:',
        '   pip install -r requirements.txt'
    ]
    
    for step in install_steps:
        if step:
            doc.add_paragraph(step)
    
    add_heading_custom(doc, 'Running Each Task', 2)
    
    doc.add_paragraph('Day 1 - Data Analysis:', style='Heading 3')
    doc.add_paragraph('python Day1/data_analysis.py')
    
    doc.add_paragraph('Day 2 - Calculator:', style='Heading 3')
    doc.add_paragraph('python Day2/calculator.py')
    
    doc.add_paragraph('Day 2 - Guessing Game:', style='Heading 3')
    doc.add_paragraph('python Day2/number_guessing_game.py')
    
    doc.add_paragraph('Day 3 - Student Manager:', style='Heading 3')
    doc.add_paragraph('python Day3/student_manager.py')
    
    doc.add_paragraph('Day 4 - Bank Demo:', style='Heading 3')
    doc.add_paragraph('python Day4/bank_demo.py')
    
    doc.add_paragraph('Day 4 - Library Demo:', style='Heading 3')
    doc.add_paragraph('python Day4/library_demo.py')
    
    doc.add_page_break()
    
    # Conclusion
    add_heading_custom(doc, 'Conclusion', 1)
    doc.add_paragraph(
        'This repository demonstrates proficiency in fundamental Python programming concepts including '
        'data analysis, control flow, file handling, and object-oriented programming. Each day builds '
        'upon previous knowledge, creating a comprehensive learning progression from basic data manipulation '
        'to advanced software design patterns.'
    )
    
    doc.add_paragraph(
        'All code follows best practices including error handling, input validation, documentation, '
        'and modular design. The repository is maintained with proper version control using Git, with '
        'daily commits following the project guidelines.'
    )
    
    # Save document
    doc.save('System_7_Repository_Documentation.docx')
    print("\n✓ Documentation created successfully!")
    print("File: System_7_Repository_Documentation.docx")

if __name__ == "__main__":
    create_documentation()
