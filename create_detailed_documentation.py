"""
Create comprehensive Word documentation with definitions from the Daily Plan
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading_custom(doc, text, level=1):
    """Add a custom styled heading"""
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return heading

def add_code_block(doc, code):
    """Add a code block with formatting"""
    para = doc.add_paragraph()
    para.style = 'No Spacing'
    run = para.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return para

def create_documentation():
    """Create the comprehensive documentation"""
    doc = Document()
    
    # Title page
    title = doc.add_heading('System 7 Repository', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Complete Documentation with Definitions & Implementation')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(0, 102, 204)
    
    author = doc.add_paragraph('Usman Bilal')
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    date = doc.add_paragraph('October 2026')
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # ========== DAY 1 ==========
    add_heading_custom(doc, 'Day 1: AI, ML and Deep Learning Foundations', 1)
    doc.add_paragraph('Commit: day-08-ml-foundations', style='Intense Quote')
    
    add_heading_custom(doc, 'Topics Covered', 2)
    doc.add_paragraph('What is AI, ML and Deep Learning; supervised, unsupervised, reinforcement learning; dataset terms')
    
    add_heading_custom(doc, 'Definitions and Concepts', 2)
    
    # AI Definition
    doc.add_paragraph('Artificial Intelligence (AI):', style='Heading 3')
    doc.add_paragraph(
        'AI is the simulation of human intelligence in machines that are programmed to think like humans '
        'and mimic their actions. AI systems can perform tasks such as learning, reasoning, problem-solving, '
        'perception, and language understanding.'
    )
    
    doc.add_paragraph('Types of AI:', style='List Bullet')
    ai_types = [
        'Narrow AI (Weak AI): Designed for specific tasks (e.g., voice assistants, recommendation systems)',
        'General AI (Strong AI): Theoretical AI with human-like intelligence across all domains',
        'Super AI: Hypothetical AI that surpasses human intelligence'
    ]
    for ai_type in ai_types:
        doc.add_paragraph(ai_type, style='List Bullet 2')
    
    # ML Definition
    doc.add_paragraph('Machine Learning (ML):', style='Heading 3')
    doc.add_paragraph(
        'Machine Learning is a subset of AI that enables systems to learn and improve from experience '
        'without being explicitly programmed. ML algorithms build mathematical models based on sample data '
        '(training data) to make predictions or decisions.'
    )
    
    # Deep Learning Definition
    doc.add_paragraph('Deep Learning:', style='Heading 3')
    doc.add_paragraph(
        'Deep Learning is a subset of Machine Learning that uses artificial neural networks with multiple '
        'layers (deep neural networks) to model complex patterns in data. It is particularly effective for '
        'processing images, speech, and text.'
    )
    
    add_heading_custom(doc, 'Learning Paradigms', 2)
    
    # Supervised Learning
    doc.add_paragraph('1. Supervised Learning:', style='Heading 3')
    doc.add_paragraph(
        'Learning from labeled data where the algorithm learns to map inputs to known outputs. '
        'The model is trained on input-output pairs and learns to predict outputs for new inputs.'
    )
    doc.add_paragraph('Examples:', style='List Bullet')
    supervised_examples = [
        'Classification: Email spam detection, image recognition',
        'Regression: House price prediction, stock price forecasting',
        'Common Algorithms: Linear Regression, Logistic Regression, Decision Trees, Random Forest, SVM, Neural Networks'
    ]
    for ex in supervised_examples:
        doc.add_paragraph(ex, style='List Bullet 2')
    
    # Unsupervised Learning
    doc.add_paragraph('2. Unsupervised Learning:', style='Heading 3')
    doc.add_paragraph(
        'Learning from unlabeled data where the algorithm tries to find hidden patterns or structures. '
        'The model explores the data without predefined labels or outcomes.'
    )
    doc.add_paragraph('Examples:', style='List Bullet')
    unsupervised_examples = [
        'Clustering: Customer segmentation, document grouping',
        'Dimensionality Reduction: Feature extraction, data compression',
        'Anomaly Detection: Fraud detection, outlier identification',
        'Common Algorithms: K-Means, Hierarchical Clustering, PCA, t-SNE, DBSCAN'
    ]
    for ex in unsupervised_examples:
        doc.add_paragraph(ex, style='List Bullet 2')
    
    # Reinforcement Learning
    doc.add_paragraph('3. Reinforcement Learning:', style='Heading 3')
    doc.add_paragraph(
        'Learning through interaction with an environment where an agent learns to make decisions by '
        'performing actions and receiving rewards or penalties. The goal is to maximize cumulative reward.'
    )
    doc.add_paragraph('Key Concepts:', style='List Bullet')
    rl_concepts = [
        'Agent: The learner or decision maker',
        'Environment: The world the agent interacts with',
        'State: Current situation of the agent',
        'Action: Decisions the agent can make',
        'Reward: Feedback from the environment',
        'Policy: Strategy that the agent follows'
    ]
    for concept in rl_concepts:
        doc.add_paragraph(concept, style='List Bullet 2')
    doc.add_paragraph('Examples: Game AI (Chess, Go), Robotics, Autonomous vehicles, Resource management')
    
    add_heading_custom(doc, 'Dataset Terms', 2)
    
    dataset_terms = {
        'Dataset': 'A collection of data used for training, testing, or validation of ML models.',
        'Features (Input Variables)': 'Individual measurable properties or characteristics of the data (e.g., age, height, income).',
        'Labels (Target Variables)': 'The output or outcome we want to predict in supervised learning.',
        'Training Set': 'Portion of data used to train the model (typically 60-80% of total data).',
        'Validation Set': 'Portion used to tune model parameters and prevent overfitting (typically 10-20%).',
        'Test Set': 'Portion used to evaluate final model performance on unseen data (typically 10-20%).',
        'Sample/Instance/Observation': 'A single row or data point in the dataset.',
        'Overfitting': 'When a model learns training data too well, including noise, and performs poorly on new data.',
        'Underfitting': 'When a model is too simple to capture the underlying patterns in the data.',
        'Bias': 'Error from incorrect assumptions in the learning algorithm.',
        'Variance': 'Error from sensitivity to small fluctuations in the training set.'
    }
    
    for term, definition in dataset_terms.items():
        para = doc.add_paragraph()
        para.add_run(f'{term}: ').bold = True
        para.add_run(definition)
    
    add_heading_custom(doc, 'Implementation: CSV Data Analysis', 2)
    doc.add_paragraph('Task: Analyze employee dataset and produce charts and summary statistics')
    doc.add_paragraph('Skills: Pandas, Matplotlib, Seaborn')
    doc.add_paragraph('Files: sample_data.csv, data_analysis.py, analysis_notebook.ipynb')
    
    doc.add_page_break()
    
    # ========== DAY 2 ==========
    add_heading_custom(doc, 'Day 2: Descriptive Statistics', 1)
    doc.add_paragraph('Commit: day-09-descriptive-statistics', style='Intense Quote')
    
    add_heading_custom(doc, 'Topics Covered', 2)
    doc.add_paragraph('Mean, median, mode, variance, standard deviation')
    
    add_heading_custom(doc, 'Definitions and Concepts', 2)
    
    # Mean
    doc.add_paragraph('Mean (Average):', style='Heading 3')
    doc.add_paragraph(
        'The arithmetic average of a dataset, calculated by summing all values and dividing by the number of values.'
    )
    doc.add_paragraph('Formula: Mean = (Sum of all values) / (Number of values)')
    doc.add_paragraph('Example: For data [2, 4, 6, 8, 10], Mean = (2+4+6+8+10)/5 = 30/5 = 6')
    doc.add_paragraph('Use Case: Represents the central tendency of data; sensitive to outliers.')
    
    # Median
    doc.add_paragraph('Median:', style='Heading 3')
    doc.add_paragraph(
        'The middle value when data is arranged in ascending or descending order. '
        'If there are even number of values, median is the average of the two middle values.'
    )
    doc.add_paragraph('Example: For [2, 4, 6, 8, 10], Median = 6 (middle value)')
    doc.add_paragraph('Example: For [2, 4, 6, 8], Median = (4+6)/2 = 5')
    doc.add_paragraph('Use Case: Better than mean when data has outliers; represents the 50th percentile.')
    
    # Mode
    doc.add_paragraph('Mode:', style='Heading 3')
    doc.add_paragraph(
        'The value that appears most frequently in a dataset. A dataset can have one mode (unimodal), '
        'two modes (bimodal), multiple modes (multimodal), or no mode.'
    )
    doc.add_paragraph('Example: For [2, 4, 4, 6, 8, 4, 10], Mode = 4 (appears 3 times)')
    doc.add_paragraph('Use Case: Useful for categorical data and finding the most common value.')
    
    # Variance
    doc.add_paragraph('Variance:', style='Heading 3')
    doc.add_paragraph(
        'A measure of how spread out the data is from the mean. It calculates the average of squared '
        'differences from the mean.'
    )
    doc.add_paragraph('Formula: Variance = Σ(xi - mean)² / n')
    doc.add_paragraph('Where xi = each value, mean = average, n = number of values')
    doc.add_paragraph('Example: For [2, 4, 6, 8, 10], mean=6')
    doc.add_paragraph('Variance = [(2-6)² + (4-6)² + (6-6)² + (8-6)² + (10-6)²] / 5')
    doc.add_paragraph('         = [16 + 4 + 0 + 4 + 16] / 5 = 40/5 = 8')
    doc.add_paragraph('Use Case: Measures data dispersion; larger variance means more spread.')
    
    # Standard Deviation
    doc.add_paragraph('Standard Deviation (SD):', style='Heading 3')
    doc.add_paragraph(
        'The square root of variance. It measures the average distance of data points from the mean '
        'and is expressed in the same units as the original data.'
    )
    doc.add_paragraph('Formula: SD = √Variance')
    doc.add_paragraph('Example: For variance = 8, SD = √8 ≈ 2.83')
    doc.add_paragraph('Use Case: Easier to interpret than variance; commonly used to measure data spread.')
    
    add_heading_custom(doc, 'Comparison Table', 2)
    doc.add_paragraph('Mean vs Median vs Mode:')
    comparison = [
        'Mean: Best for symmetric data without outliers',
        'Median: Best for skewed data or data with outliers',
        'Mode: Best for categorical data or finding most frequent value',
        'Variance & SD: Measure how much data varies from the mean'
    ]
    for item in comparison:
        doc.add_paragraph(item, style='List Bullet')
    
    add_heading_custom(doc, 'Implementation: Calculator & Guessing Game', 2)
    doc.add_paragraph('Task: Build command-line calculator and number guessing game')
    doc.add_paragraph('Skills: Input/output, loops, functions, control flow')
    doc.add_paragraph('Files: calculator.py, number_guessing_game.py')
    
    doc.add_page_break()
    
    # ========== DAY 3 ==========
    add_heading_custom(doc, 'Day 3: Linear Algebra Basics', 1)
    doc.add_paragraph('Commit: day-10-linear-algebra', style='Intense Quote')
    
    add_heading_custom(doc, 'Topics Covered', 2)
    doc.add_paragraph('Vectors, matrices, dot products')
    
    add_heading_custom(doc, 'Definitions and Concepts', 2)
    
    # Vectors
    doc.add_paragraph('Vectors:', style='Heading 3')
    doc.add_paragraph(
        'A vector is an ordered array of numbers representing magnitude and direction. '
        'In machine learning, vectors represent features, data points, or weights.'
    )
    doc.add_paragraph('Types:', style='List Bullet')
    vector_types = [
        'Row Vector: [1, 2, 3] - horizontal arrangement',
        'Column Vector: [[1], [2], [3]] - vertical arrangement',
        'Example: A data point with features [age=25, height=170, weight=70]'
    ]
    for vtype in vector_types:
        doc.add_paragraph(vtype, style='List Bullet 2')
    
    doc.add_paragraph('Vector Operations:', style='List Bullet')
    vector_ops = [
        'Addition: [1,2] + [3,4] = [4,6]',
        'Scalar Multiplication: 2 × [1,2,3] = [2,4,6]',
        'Magnitude (Length): ||v|| = √(x² + y² + z²)',
        'Example: ||[3,4]|| = √(9+16) = √25 = 5'
    ]
    for op in vector_ops:
        doc.add_paragraph(op, style='List Bullet 2')
    
    # Matrices
    doc.add_paragraph('Matrices:', style='Heading 3')
    doc.add_paragraph(
        'A matrix is a 2D array of numbers arranged in rows and columns. In ML, matrices represent '
        'datasets, transformation operations, and neural network weights.'
    )
    
    matrix_example = '''Example 3×3 Matrix:
    [1  2  3]
    [4  5  6]
    [7  8  9]
    
Dimensions: m × n (m rows, n columns)'''
    add_code_block(doc, matrix_example)
    
    doc.add_paragraph('Matrix Operations:', style='List Bullet')
    matrix_ops = [
        'Addition: Add corresponding elements (matrices must have same dimensions)',
        'Multiplication: Multiply rows by columns (inner dimensions must match)',
        'Transpose: Flip rows and columns (AT)',
        'Identity Matrix: Square matrix with 1s on diagonal, 0s elsewhere',
        'Inverse Matrix: A × A⁻¹ = I (used in solving equations)'
    ]
    for op in matrix_ops:
        doc.add_paragraph(op, style='List Bullet 2')
    
    # Dot Product
    doc.add_paragraph('Dot Product (Scalar Product):', style='Heading 3')
    doc.add_paragraph(
        'The dot product of two vectors is the sum of the products of their corresponding elements. '
        'It results in a scalar (single number).'
    )
    doc.add_paragraph('Formula: A · B = a₁b₁ + a₂b₂ + a₃b₃ + ... + aₙbₙ')
    
    dot_example = '''Example:
A = [1, 2, 3]
B = [4, 5, 6]

A · B = (1×4) + (2×5) + (3×6)
      = 4 + 10 + 18
      = 32'''
    add_code_block(doc, dot_example)
    
    doc.add_paragraph('Properties and Uses:', style='List Bullet')
    dot_uses = [
        'Measures similarity between vectors (higher value = more similar)',
        'Used in neural networks for weighted sums',
        'Used in calculating projections and angles',
        'Cosine Similarity: cos(θ) = (A·B) / (||A|| × ||B||)',
        'If dot product = 0, vectors are perpendicular (orthogonal)'
    ]
    for use in dot_uses:
        doc.add_paragraph(use, style='List Bullet 2')
    
    add_heading_custom(doc, 'ML Applications', 2)
    doc.add_paragraph('How Linear Algebra is Used in Machine Learning:', style='List Bullet')
    ml_apps = [
        'Feature Representation: Each data point is a vector of features',
        'Dataset Representation: Entire dataset is a matrix (rows=samples, columns=features)',
        'Neural Networks: Weights are matrices, forward propagation uses matrix multiplication',
        'PCA (Dimensionality Reduction): Uses eigenvectors and eigenvalues',
        'Gradient Descent: Updates weights using vector operations',
        'Image Processing: Images are matrices of pixel values'
    ]
    for app in ml_apps:
        doc.add_paragraph(app, style='List Bullet 2')
    
    add_heading_custom(doc, 'Implementation: Student Record Manager', 2)
    doc.add_paragraph('Task: Create student record manager with file handling')
    doc.add_paragraph('Skills: Dictionaries, file handling (JSON), exceptions')
    doc.add_paragraph('Files: student_manager.py, students.json')
    
    doc.add_page_break()
    
    # ========== DAY 4 ==========
    add_heading_custom(doc, 'Day 4: Probability & Data Preparation', 1)
    doc.add_paragraph('Commit: day-11-probability-data-prep', style='Intense Quote')
    
    add_heading_custom(doc, 'Topics Covered', 2)
    doc.add_paragraph('Probability, correlation, train/validation/test split; dataset exploration task')
    
    add_heading_custom(doc, 'Definitions and Concepts', 2)
    
    # Probability
    doc.add_paragraph('Probability:', style='Heading 3')
    doc.add_paragraph(
        'Probability is the measure of the likelihood that an event will occur. '
        'It ranges from 0 (impossible) to 1 (certain).'
    )
    doc.add_paragraph('Formula: P(Event) = (Number of favorable outcomes) / (Total number of outcomes)')
    doc.add_paragraph('Example: Probability of rolling a 6 on a die = 1/6 ≈ 0.167')
    
    doc.add_paragraph('Key Probability Concepts:', style='List Bullet')
    prob_concepts = [
        'Independent Events: P(A and B) = P(A) × P(B)',
        'Dependent Events: P(A and B) = P(A) × P(B|A)',
        'Conditional Probability: P(A|B) = P(A and B) / P(B)',
        'Bayes\' Theorem: P(A|B) = [P(B|A) × P(A)] / P(B)',
        'Example: Probability of spam email given certain words appear'
    ]
    for concept in prob_concepts:
        doc.add_paragraph(concept, style='List Bullet 2')
    
    # Correlation
    doc.add_paragraph('Correlation:', style='Heading 3')
    doc.add_paragraph(
        'Correlation measures the strength and direction of the relationship between two variables. '
        'It ranges from -1 to +1.'
    )
    
    doc.add_paragraph('Correlation Values:', style='List Bullet')
    corr_values = [
        '+1: Perfect positive correlation (both increase together)',
        '0: No correlation (no relationship)',
        '-1: Perfect negative correlation (one increases, other decreases)',
        'Example: +0.88 correlation between years of experience and salary (strong positive)'
    ]
    for val in corr_values:
        doc.add_paragraph(val, style='List Bullet 2')
    
    doc.add_paragraph('Correlation vs Causation:', style='Heading 4')
    doc.add_paragraph(
        'Important: Correlation does NOT imply causation! '
        'Two variables may be correlated without one causing the other.'
    )
    doc.add_paragraph('Example: Ice cream sales and drowning deaths are correlated (both increase in summer), '
                     'but ice cream doesn\'t cause drowning!')
    
    # Train/Validation/Test Split
    doc.add_paragraph('Train/Validation/Test Split:', style='Heading 3')
    doc.add_paragraph(
        'Dividing dataset into separate portions to train, tune, and evaluate machine learning models.'
    )
    
    doc.add_paragraph('Training Set (60-80%):', style='Heading 4')
    doc.add_paragraph(
        'Data used to train the model - the model learns patterns from this data. '
        'This is the largest portion of your data.'
    )
    
    doc.add_paragraph('Validation Set (10-20%):', style='Heading 4')
    doc.add_paragraph(
        'Data used to tune hyperparameters and prevent overfitting during training. '
        'Helps you choose the best model configuration.'
    )
    
    doc.add_paragraph('Test Set (10-20%):', style='Heading 4')
    doc.add_paragraph(
        'Data used ONLY ONCE at the end to evaluate final model performance. '
        'Simulates real-world unseen data. Never use this for training or tuning!'
    )
    
    split_example = '''Example Split (100 samples):
Training:   60 samples (60%)
Validation: 20 samples (20%)
Test:       20 samples (20%)

Common Splits:
- 80/10/10
- 70/15/15
- 60/20/20'''
    add_code_block(doc, split_example)
    
    doc.add_paragraph('Why Split Data?', style='List Bullet')
    why_split = [
        'Prevents overfitting to training data',
        'Provides unbiased evaluation of model',
        'Allows tuning without "cheating" by seeing test data',
        'Simulates model performance on new, unseen data',
        'Helps detect if model generalizes well'
    ]
    for reason in why_split:
        doc.add_paragraph(reason, style='List Bullet 2')
    
    # Dataset Exploration
    doc.add_paragraph('Dataset Exploration:', style='Heading 3')
    doc.add_paragraph('Key steps in exploring a dataset before training models:')
    
    exploration_steps = [
        'Load and inspect data (head, tail, shape, info)',
        'Check data types and missing values',
        'Calculate summary statistics (mean, median, std)',
        'Visualize distributions (histograms, box plots)',
        'Check for outliers and anomalies',
        'Analyze correlations between features',
        'Handle missing data (imputation or removal)',
        'Encode categorical variables if needed',
        'Scale/normalize features if required',
        'Split into train/validation/test sets'
    ]
    for i, step in enumerate(exploration_steps, 1):
        doc.add_paragraph(f'{i}. {step}', style='List Number')
    
    add_heading_custom(doc, 'Implementation: Bank Account & Library OOP', 2)
    doc.add_paragraph('Task: Design Bank Account and Library classes with OOP')
    doc.add_paragraph('Skills: Classes, inheritance, encapsulation, methods')
    doc.add_paragraph('Files: bank_account.py, library_system.py, bank_demo.py, library_demo.py')
    
    doc.add_paragraph('OOP Concepts Applied:', style='List Bullet')
    oop_applied = [
        'Classes: BankAccount, SavingsAccount, CheckingAccount, Library, Book, Member',
        'Inheritance: SavingsAccount and CheckingAccount inherit from BankAccount',
        'Encapsulation: Private balance attribute (_balance)',
        'Methods: deposit(), withdraw(), borrow_book(), return_book()',
        'Constructor: __init__() for initialization',
        'String representation: __str__() for readable output'
    ]
    for concept in oop_applied:
        doc.add_paragraph(concept, style='List Bullet 2')
    
    doc.add_page_break()
    
    # ========== SUMMARY ==========
    add_heading_custom(doc, 'Repository Summary', 1)
    
    add_heading_custom(doc, 'Completed Tasks', 2)
    completed = [
        'Day 1: CSV Data Analysis with Pandas, Matplotlib, Seaborn',
        'Day 2: Command-line Calculator and Number Guessing Game',
        'Day 3: Student Record Manager with File Handling',
        'Day 4: Bank Account & Library Management Systems (OOP)'
    ]
    for task in completed:
        doc.add_paragraph(task, style='List Bullet')
    
    add_heading_custom(doc, 'Skills Demonstrated', 2)
    skills = [
        'Data analysis and visualization',
        'Control flow (loops, conditionals)',
        'Functions and modular programming',
        'File I/O and JSON handling',
        'Exception handling and validation',
        'Object-Oriented Programming',
        'Git version control',
        'Documentation and code organization'
    ]
    for skill in skills:
        doc.add_paragraph(skill, style='List Bullet')
    
    add_heading_custom(doc, 'GitHub Repository', 2)
    doc.add_paragraph('https://github.com/Usman-Bilal/System-7-Repository')
    
    add_heading_custom(doc, 'How to Run', 2)
    doc.add_paragraph('1. Clone repository: git clone https://github.com/Usman-Bilal/System-7-Repository.git')
    doc.add_paragraph('2. Install requirements: pip install -r requirements.txt')
    doc.add_paragraph('3. Run any day\'s code:')
    doc.add_paragraph('   - python Day1/data_analysis.py', style='List Bullet')
    doc.add_paragraph('   - python Day2/calculator.py', style='List Bullet')
    doc.add_paragraph('   - python Day3/student_manager.py', style='List Bullet')
    doc.add_paragraph('   - python Day4/bank_demo.py', style='List Bullet')
    
    # Save document
    doc.save('System_7_Complete_Documentation.docx')
    print("\n✓ Complete documentation created successfully!")
    print("File: System_7_Complete_Documentation.docx")

if __name__ == "__main__":
    create_documentation()
